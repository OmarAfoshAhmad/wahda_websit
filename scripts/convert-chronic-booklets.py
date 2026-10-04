"""
يحوّل كتيّبات الأدوية المزمنة (كتلة لكل مستفيد: الاسم، رقم البطاقة، جدول الأدوية) إلى ملف
استيراد مسطح بالبنية التي تقرؤها صفحة "استيراد الأدوية المزمنة":
رقم البطاقة | اسم المستفيد | اسم الدواء | الجرعة | التكرار | ملاحظات

الاستخدام:
    python scripts/convert-chronic-booklets.py "<مجلد الكتيبات>" "<مجلد الإخراج>"

كل ملف كتيّب يُنتج ملف استيراد بنفس الاسم، وورقة "مراجعة" بالكتل التي تحتاج تدقيقًا
(بلا بطاقة، أو بطاقة غير قياسية، أو بلا أدوية). والملخص يُطبع في النهاية.
"""
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

HEADERS = ["رقم البطاقة", "اسم المستفيد", "اسم الدواء", "الجرعة", "التكرار", "ملاحظات"]
EMPTY_MARKERS = {"", "/", "-", "*", "**", "***", "0", "none"}


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def clean(value) -> str:
    result = text(value)
    return "" if result.lower() in EMPTY_MARKERS else result


def letters_only(value: str) -> str:
    return re.sub(r"[^؀-ۿa-zA-Z]", "", value)


def is_name_label(value) -> bool:
    label = letters_only(text(value)).replace("ـ", "")
    return label in {"الاسم", "الأسم", "الإسم"}


def is_card_label(value) -> bool:
    return "البطاقة" in text(value) or "البطاقه" in text(value)


def first_value_right(row: tuple, column: int, span: int = 4) -> str:
    for offset in range(1, span + 1):
        if column + offset < len(row):
            value = clean(row[column + offset])
            if value:
                return value
    return ""


def card_is_standard(card: str) -> bool:
    # البطاقات القياسية تبدأ بحروف ثم أرقام (WAAD2025..., JMR2025..., LCC2025...).
    return bool(re.match(r"^[A-Za-z][A-Za-z0-9]+\d{4,}", card))


def extract_blocks(sheet: Worksheet):
    rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
    blocks = []
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            if not is_name_label(value):
                continue
            name = first_value_right(row, c)
            card = ""
            header_row = None
            for look in range(r + 1, min(r + 5, len(rows))):
                candidate = rows[look]
                if c < len(candidate) and is_card_label(candidate[c]):
                    card = first_value_right(candidate, c).upper().replace(" ", "")
                if c < len(candidate) and text(candidate[c]) == "ت":
                    header_row = look
                    break
            drugs = []
            if header_row is not None:
                for d in range(header_row + 1, len(rows)):
                    line = rows[d]
                    if c < len(line) and (is_name_label(line[c]) or is_card_label(line[c])):
                        break
                    cell = lambda k: clean(line[c + k]) if c + k < len(line) else ""
                    drug = cell(1)
                    if drug:
                        drugs.append((drug, cell(2), cell(3), cell(4)))
            blocks.append({"sheet": sheet.title, "row": r + 1, "name": name, "card": card, "drugs": drugs})
    return blocks


def convert(source: Path, target_dir: Path):
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    blocks = [block for sheet in workbook.worksheets for block in extract_blocks(sheet)]

    out = openpyxl.Workbook()
    data = out.active
    data.title = "الأدوية المزمنة"
    data.sheet_view.rightToLeft = True
    data.append(HEADERS)
    review = out.create_sheet("مراجعة")
    review.sheet_view.rightToLeft = True
    review.append(["الورقة", "الصف", "اسم المستفيد", "رقم البطاقة", "عدد الأدوية", "السبب"])

    stats = {"blocks": 0, "rows": 0, "review": 0, "empty": 0}
    for block in blocks:
        if not block["name"] and not block["card"] and not block["drugs"]:
            stats["empty"] += 1  # كتلة قالب فارغة
            continue
        stats["blocks"] += 1
        reasons = []
        if not block["drugs"]:
            reasons.append("لا توجد أدوية")
        if not block["card"]:
            reasons.append("لا يوجد رقم بطاقة: سيُطابق بالاسم")
        elif not card_is_standard(block["card"]):
            reasons.append("رقم بطاقة غير قياسي: قد يُطابق بالاسم")
        if not block["name"]:
            reasons.append("لا يوجد اسم")
        if reasons:
            stats["review"] += 1
            review.append([block["sheet"], block["row"], block["name"], block["card"], len(block["drugs"]), "، ".join(reasons)])
        for drug, dose, frequency, notes in block["drugs"]:
            data.append([block["card"], block["name"], drug, dose, frequency, notes])
            stats["rows"] += 1

    for sheet, widths in ((data, [22, 34, 30, 16, 14, 30]), (review, [14, 8, 34, 22, 12, 50])):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for index, width in enumerate(widths):
            sheet.column_dimensions[chr(65 + index)].width = width
        sheet.freeze_panes = "A2"

    target = target_dir / f"استيراد - {source.stem}.xlsx"
    out.save(target)
    return target, stats


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    source_dir, target_dir = Path(sys.argv[1]), Path(sys.argv[2])
    target_dir.mkdir(parents=True, exist_ok=True)
    for source in sorted(source_dir.glob("*.xlsx")):
        if source.name.startswith("~$"):
            continue
        target, stats = convert(source, target_dir)
        print(f"{source.name}: {stats['blocks']} مستفيد، {stats['rows']} صف دواء، {stats['review']} للمراجعة، {stats['empty']} كتلة فارغة -> {target.name}")


if __name__ == "__main__":
    main()
